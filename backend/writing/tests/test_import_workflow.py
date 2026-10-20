"""End-to-end importer coverage through the published student catalogue."""

import hashlib
import secrets
import tempfile
from datetime import timedelta
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from accounts.models import StudentSession
from content.models import LearningConfiguration
from entitlements.services import grant_entitlement
from writing.import_files import make_template, parse_workbook
from writing.import_services import build_import_preview, confirm_import
from writing.import_storage import load_preview, save_preview
from writing.models import WritingContent, WritingRevision
from writing.services import make_publication_token, publish_writings


@override_settings(ROOT_URLCONF="writing.tests.urls")
class WritingImportWorkflowTests(TestCase):
    def setUp(self):
        self.private = tempfile.TemporaryDirectory()
        self.addCleanup(self.private.cleanup)
        settings = override_settings(WRITING_IMPORT_TMP_DIR=self.private.name)
        settings.enable()
        self.addCleanup(settings.disable)
        self.operator = get_user_model().objects.create_superuser(
            username="import-workflow-operator", password="test-only-import"
        )
        self.student = get_user_model().objects.create_user(username="import-workflow-student")
        self.bearer = secrets.token_urlsafe(32)
        StudentSession.objects.create(
            user=self.student,
            token_digest=hashlib.sha256(self.bearer.encode()).hexdigest(),
            encrypted_session_key="test-only-ciphertext",
            expires_at=timezone.now() + timedelta(days=1),
        )
        LearningConfiguration.objects.create(page_size=2)
        self.client = APIClient()
        self.rows = [
            ["import-00", "import-category", "import-genre", "写作模板", "模板",
             "  <b>通知格式</b>\n\n    段落缩进保留", "", 0],
            ["import-01", "import-category", "import-genre", "范文示例", "范文", "范文正文", "", 1],
            ["import-02", "import-category", "import-genre", "历年真题", "真题", "真题正文", "", 2],
            ["import-03", "import-category", "import-genre", "专项练习", "练习", "练习题干", "答案\n  第二行", 3],
        ]

    def workbook(self, rows=None):
        workbook = load_workbook(BytesIO(make_template()))
        for row_number, values in enumerate(rows if rows is not None else self.rows, 2):
            for column, value in enumerate(values, 1):
                workbook["应用文"].cell(row_number, column, value)
        for sheet, values in (
            ("分类", ["import-category", "导入测试分类", 10]),
            ("文种", ["import-genre", "通知", 10]),
        ):
            for column, value in enumerate(values, 1):
                workbook[sheet].cell(2, column, value)
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        return stream.getvalue()

    def import_drafts(self, rows=None, mode="create_only"):
        raw = self.workbook(rows)
        parsed = parse_workbook(raw, "应用文.xlsx")
        self.assertEqual(parsed["issues"], [])
        preview = build_import_preview(parsed, self.operator, mode)
        self.assertTrue(preview["canConfirm"], preview["issues"])
        saved = save_preview(self.operator, raw, preview)
        loaded = load_preview(saved["previewId"], self.operator, saved["confirmationToken"])
        return confirm_import(loaded, self.operator)

    def publish(self):
        contents = WritingContent.objects.all()
        token = make_publication_token(list(contents), self.operator, "publish_selected")
        publish_writings(contents, self.operator, confirmation_token=token)

    def test_import_publish_and_shared_free_set_with_four_kinds(self):
        self.import_drafts()
        self.assertEqual(WritingContent.objects.count(), 4)
        self.assertFalse(WritingContent.objects.filter(is_published=True).exists())
        self.assertEqual(WritingRevision.objects.count(), 0)
        self.assertEqual(self.client.get("/api/student/v1/writings/import-00/").status_code, 404)
        self.publish()
        response = self.client.get("/api/student/v1/writings/?categoryId=import-category")
        self.assertEqual(response.status_code, 200)
        self.assertEqual([item["id"] for item in response.data["results"]], ["import-00", "import-01"])
        self.assertEqual(response.data["access"]["restrictedCount"], 2)
        for kind in ("past_exam", "exercise"):
            filtered = self.client.get("/api/student/v1/writings/?categoryId=import-category&kind=" + kind)
            self.assertEqual(filtered.status_code, 200)
            self.assertEqual(filtered.data["results"], [])
        paid = self.client.get("/api/student/v1/writings/import-03/")
        self.assertEqual(paid.status_code, 403)
        self.assertNotIn("练习题干", str(paid.data))
        self.assertNotIn("第二行", str(paid.data))
        free = self.client.get("/api/student/v1/writings/import-00/")
        self.assertEqual(free.data["body"], self.rows[0][5])
        self.client.credentials(HTTP_AUTHORIZATION="Bearer " + self.bearer)
        grant_entitlement(self.student, "activation", "import-workflow", actor=self.operator)
        exercise = self.client.get("/api/student/v1/writings/import-03/")
        self.assertEqual(exercise.status_code, 200)
        self.assertEqual(exercise.data["referenceAnswer"], self.rows[3][6])
        second_page = self.client.get("/api/student/v1/writings/?categoryId=import-category&page=2")
        self.assertEqual(second_page.status_code, 200)
        self.assertEqual([item["id"] for item in second_page.data["results"]], ["import-02", "import-03"])

    def test_imported_draft_update_keeps_published_snapshot_until_republish(self):
        self.import_drafts()
        self.publish()
        previous = WritingContent.objects.get(pk="import-00").published_revision_id
        updated = [list(self.rows[0])]
        updated[0][5] = "新版正文\n  新版缩进"
        self.import_drafts(updated, mode="upsert_draft")
        content = WritingContent.objects.get(pk="import-00")
        self.assertTrue(content.is_published)
        self.assertEqual(content.published_revision_id, previous)
        self.assertEqual(content.body, updated[0][5])
        self.assertEqual(self.client.get("/api/student/v1/writings/import-00/").data["body"], self.rows[0][5])
        self.publish()
        self.assertEqual(self.client.get("/api/student/v1/writings/import-00/").data["body"], updated[0][5])
        self.assertEqual(WritingRevision.objects.filter(content_id="import-00").count(), 2)
        self.assertEqual(WritingContent.objects.count(), 4)
