"""Workbook contract, normalization and hostile-package regression tests."""

from datetime import date
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from django.test import SimpleTestCase
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.protection import WorkbookProtection

from writing.import_files import (
    CATEGORY_HEADERS, CONTENT_HEADERS, GENRE_HEADERS, MAX_EXPANDED_BYTES,
    MAX_UPLOAD_BYTES, make_error_report, make_template, parse_workbook,
)


def save_workbook(workbook):
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def content_row(**changes):
    row = {
        "内容编号": "notice-001", "分类编号": "daily-writing", "文种编号": "notice",
        "内容类型": "写作模板", "标题": "通知写作格式", "正文": "正文\n  缩进段落",
        "参考答案": None, "排序": 10,
    }
    row.update(changes)
    return [row[header] for header in CONTENT_HEADERS]


def build_workbook(rows=None, *, dictionaries=True):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "应用文"
    worksheet.append(CONTENT_HEADERS)
    for row in [content_row()] if rows is None else rows:
        worksheet.append(row)
    if dictionaries:
        category = workbook.create_sheet("分类")
        category.append(CATEGORY_HEADERS)
        category.append(["daily-writing", "日常应用文", 10])
        genre = workbook.create_sheet("文种")
        genre.append(GENRE_HEADERS)
        genre.append(["notice", "通知", 20])
    return workbook


def change_package(raw, replacements=None, additions=None):
    replacements = replacements or {}
    output = BytesIO()
    with ZipFile(BytesIO(raw)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            value = source.read(entry)
            change = replacements.get(entry.filename)
            target.writestr(entry.filename, change(value) if callable(change) else change if change is not None else value)
        for name, value in (additions or {}).items():
            target.writestr(name, value)
    return output.getvalue()


class WorkbookImportTests(SimpleTestCase):
    def parse(self, workbook, filename="应用文.xlsx"):
        return parse_workbook(save_workbook(workbook), filename)

    def assert_issue(self, result, *, sheet=None, row=None, column=None, code=None):
        matches = [issue for issue in result["issues"] if
                   (sheet is None or issue["sheet"] == sheet)
                   and (row is None or issue["row"] == row)
                   and (column is None or issue["column"] == column)
                   and (code is None or issue["code"] == code)]
        self.assertTrue(matches, result["issues"])
        for issue in result["issues"]:
            self.assertEqual(set(issue), {"sheet", "row", "column", "code", "message"})

    def test_four_kinds_and_normalized_model_fields(self):
        rows = [content_row(**{"内容编号": f"id-{index}", "内容类型": kind,
                              "参考答案": "答案" if index == 3 else None})
                for index, kind in enumerate(("写作模板", "范文示例", "past_exam", " exercise "))]
        result = self.parse(build_workbook(rows))
        self.assertEqual(result["issues"], [])
        self.assertEqual([item["values"]["kind"] for item in result["contents"]],
                         ["template", "example", "past_exam", "exercise"])
        self.assertEqual(result["contents"][0], {"row": 2, "values": {
            "id": "id-0", "category_id": "daily-writing", "genre_id": "notice", "kind": "template",
            "title": "通知写作格式", "body": "正文\n  缩进段落", "reference_answer": "", "sort_order": 10,
        }})
        self.assertEqual(result["categories"], [{"row": 2, "values": {"id": "daily-writing", "name": "日常应用文", "sort_order": 10}}])
        self.assertEqual(result["genres"][0]["values"]["id"], "notice")

    def test_trim_only_ids_kind_title_name_and_preserve_body(self):
        body = "  首段\r\n\r\n\t次段\r末尾  \n字面\\n和↵"
        answer = "\t答案\r\n  第二段\r"
        workbook = build_workbook([content_row(**{
            "内容编号": " 001 ", "分类编号": " daily-writing ", "文种编号": " notice ",
            "内容类型": " 专项练习 ", "标题": " 标题 ", "正文": body, "参考答案": answer,
        })])
        workbook["分类"]["B2"] = " 日常应用文 "
        # On Windows openpyxl's stdlib writer translates literal CRLF to CRCRLF.
        # Encode CR explicitly so this fixture represents real uploaded CRLF.
        raw = change_package(save_workbook(workbook), {
            "xl/worksheets/sheet1.xml": lambda value: value.replace(b"\r\r\n", b"&#13;\n"),
        })
        result = parse_workbook(raw, "file.xlsx")
        self.assertEqual(result["issues"], [])
        values = result["contents"][0]["values"]
        self.assertEqual(values["id"], "001")
        self.assertEqual(values["title"], "标题")
        self.assertEqual(values["body"], body.replace("\r\n", "\n").replace("\r", "\n"))
        self.assertEqual(values["reference_answer"], "\t答案\n  第二段\n")
        self.assertEqual(result["categories"][0]["values"]["name"], "日常应用文")

    def test_optional_sheets_and_blank_values(self):
        workbook = build_workbook([content_row(**{"文种编号": " \t", "参考答案": "\n", "排序": None})], dictionaries=False)
        result = self.parse(workbook, "template.XLSX")
        self.assertEqual(result["issues"], [])
        self.assertIsNone(result["contents"][0]["values"]["genre_id"])
        self.assertEqual(result["contents"][0]["values"]["reference_answer"], "")
        self.assertEqual(result["contents"][0]["values"]["sort_order"], 0)

    def test_headers_can_be_reordered(self):
        workbook = build_workbook(rows=[])
        workbook["应用文"].delete_rows(1)
        workbook["应用文"].append(list(reversed(CONTENT_HEADERS)))
        workbook["应用文"].append(list(reversed(content_row())))
        result = self.parse(workbook)
        self.assertEqual(result["issues"], [])
        self.assertEqual(result["contents"][0]["values"]["id"], "notice-001")

    def test_missing_duplicate_unknown_and_inexact_headers_are_rejected(self):
        for headers in (CONTENT_HEADERS[:-1], (*CONTENT_HEADERS, "价格"),
                        (*CONTENT_HEADERS[:-1], "正文"), (" 内容编号", *CONTENT_HEADERS[1:])):
            with self.subTest(headers=headers):
                workbook = build_workbook(rows=[])
                workbook["应用文"].delete_rows(1)
                workbook["应用文"].append(headers)
                self.assert_issue(self.parse(workbook), sheet="应用文", row=1, code="INVALID_HEADER")

    def test_values_in_unheaded_extra_column_are_rejected(self):
        workbook = build_workbook()
        workbook["应用文"]["I2"] = "不应导入的字段"
        self.assert_issue(self.parse(workbook), sheet="应用文", row=1, code="INVALID_HEADER")

    def test_blank_rows_and_preformatted_cells_keep_original_row_numbers(self):
        workbook = build_workbook(rows=[])
        worksheet = workbook["应用文"]
        worksheet.append([" "] * len(CONTENT_HEADERS))
        worksheet.append([None] * len(CONTENT_HEADERS))
        worksheet.append(content_row())
        worksheet["XFD1048576"].number_format = "@"
        result = self.parse(workbook)
        self.assertEqual(result["issues"], [])
        self.assertEqual(result["contents"][0]["row"], 4)

    def test_bad_dimension_does_not_drop_physical_data(self):
        raw = change_package(save_workbook(build_workbook()), {
            "xl/worksheets/sheet1.xml": lambda value: value.replace(b'<dimension ref="A1:H2"/>', b'<dimension ref="A1:A1"/>'),
        })
        result = parse_workbook(raw, "file.xlsx")
        self.assertEqual(result["issues"], [])
        self.assertEqual(len(result["contents"]), 1)

    def test_duplicate_ids_report_actual_excel_rows(self):
        workbook = build_workbook([content_row(), [], content_row()])
        workbook["分类"].append(["daily-writing", "日常应用文", 10])
        workbook["文种"].append(["notice", "通知", 20])
        result = self.parse(workbook)
        self.assert_issue(result, sheet="应用文", row=4, column="内容编号", code="DUPLICATE_ID")
        self.assert_issue(result, sheet="分类", row=3, code="DUPLICATE_ID")
        self.assert_issue(result, sheet="文种", row=3, code="DUPLICATE_ID")

    def test_ids_reject_numbers_boolean_unicode_and_overlong_values(self):
        for value in (123, 1.2, True, "汉字", "abc.1", "a" * 65, "abc\ndef"):
            for header in ("内容编号", "分类编号", "文种编号"):
                with self.subTest(value=value, header=header):
                    result = self.parse(build_workbook([content_row(**{header: value})]))
                    self.assert_issue(result, sheet="应用文", row=2, column=header, code="INVALID_VALUE")

    def test_valid_ascii_id_boundary_and_case_are_preserved(self):
        identifiers = ("A_a-" + "0" * 60, "same", "SAME")
        result = self.parse(build_workbook([content_row(**{"内容编号": identifier}) for identifier in identifiers]))
        self.assertEqual(result["issues"], [])
        self.assertEqual([item["values"]["id"] for item in result["contents"]], list(identifiers))

    def test_sort_rejects_text_boolean_fraction_date_and_out_of_range(self):
        for value in ("1", "1e2", " ", True, False, 1.5, date(2026, 10, 20), -1, 2147483648):
            with self.subTest(value=value):
                result = self.parse(build_workbook([content_row(**{"排序": value})]))
                self.assert_issue(result, row=2, column="排序", code="INVALID_VALUE")
        for value in (0, 2147483647):
            with self.subTest(value=value):
                result = self.parse(build_workbook([content_row(**{"排序": value})]))
                self.assertEqual(result["issues"], [])
                self.assertEqual(result["contents"][0]["values"]["sort_order"], value)

    def test_required_fields_and_answer_conditions(self):
        for header in ("内容编号", "分类编号", "内容类型", "标题", "正文"):
            with self.subTest(header=header):
                self.assert_issue(self.parse(build_workbook([content_row(**{header: " \n"})])), row=2, column=header, code="REQUIRED_FIELD")
        self.assert_issue(self.parse(build_workbook([content_row(**{"内容类型": "exercise"})])), row=2, column="参考答案", code="REQUIRED_FIELD")
        self.assert_issue(self.parse(build_workbook([content_row(**{"参考答案": "不可填写"})])), row=2, column="参考答案", code="INVALID_VALUE")
        self.assert_issue(self.parse(build_workbook([content_row(**{"内容类型": "通知"})])), row=2, column="内容类型", code="INVALID_VALUE")

    def test_text_length_boundaries_and_wrong_cell_types(self):
        result = self.parse(build_workbook([content_row(**{"标题": "题" * 200, "正文": "文" * 20000, "内容类型": "exercise", "参考答案": "答" * 20000})]))
        self.assertEqual(result["issues"], [])
        for header, value in (("标题", "题" * 201), ("正文", "文" * 20001), ("参考答案", "答" * 20001), ("正文", 123), ("标题", True)):
            with self.subTest(header=header):
                self.assert_issue(self.parse(build_workbook([content_row(**{header: value})])), row=2, column=header, code="INVALID_VALUE")
        workbook = build_workbook()
        workbook["分类"]["B2"] = "类" * 121
        workbook["文种"]["B2"] = "种" * 121
        result = self.parse(workbook)
        self.assert_issue(result, sheet="分类", column="分类名称", code="INVALID_VALUE")
        self.assert_issue(result, sheet="文种", column="文种名称", code="INVALID_VALUE")

    def test_empty_content_is_an_error_even_with_dictionaries(self):
        result = self.parse(build_workbook(rows=[]))
        self.assert_issue(result, sheet="应用文", code="REQUIRED_FIELD")
        self.assertEqual(result["contents"], [])

    def test_row_limits_count_nonempty_rows_and_stop_at_first_excess(self):
        workbook = build_workbook([content_row(**{"内容编号": f"id-{index}"}) for index in range(1001)])
        for index in range(200):
            workbook["分类"].append([f"category-{index}", f"分类 {index}", index])
            workbook["文种"].append([f"genre-{index}", f"文种 {index}", index])
        result = self.parse(workbook)
        self.assert_issue(result, sheet="应用文", row=1002, code="INVALID_VALUE")
        self.assert_issue(result, sheet="分类", row=202, code="INVALID_VALUE")
        self.assert_issue(result, sheet="文种", row=202, code="INVALID_VALUE")
        self.assertEqual(len(result["contents"]), 1000)
        self.assertEqual(len(result["categories"]), 200)

    def test_exact_row_limits_are_accepted(self):
        workbook = build_workbook([content_row(**{"内容编号": f"id-{index}"}) for index in range(1000)])
        for index in range(199):
            workbook["分类"].append([f"category-{index}", f"分类 {index}", index])
            workbook["文种"].append([f"genre-{index}", f"文种 {index}", index])
        result = self.parse(workbook)
        self.assertEqual(result["issues"], [])
        self.assertEqual((len(result["contents"]), len(result["categories"]), len(result["genres"])), (1000, 200, 200))

    def test_formula_and_excel_errors_are_located_with_chinese_column(self):
        for value in ("=1+1", "#DIV/0!"):
            with self.subTest(value=value):
                workbook = build_workbook()
                workbook["应用文"]["F2"] = value
                self.assert_issue(self.parse(workbook), sheet="应用文", row=2, column="正文", code="INVALID_VALUE")

    def test_formula_in_guidance_is_rejected(self):
        workbook = build_workbook()
        workbook.create_sheet("填写说明")["B7"] = "=HYPERLINK(\"https://example.com\")"
        self.assert_issue(self.parse(workbook), sheet="工作簿", row=7, code="INVALID_VALUE")

    def test_merged_hidden_rows_columns_and_protection_are_rejected(self):
        changes = (
            lambda workbook: workbook["应用文"].merge_cells("A2:B2"),
            lambda workbook: setattr(workbook["应用文"].row_dimensions[3], "hidden", True),
            lambda workbook: setattr(workbook["应用文"].column_dimensions["J"], "hidden", True),
            lambda workbook: setattr(workbook["应用文"].protection, "sheet", True),
            lambda workbook: setattr(workbook["分类"], "sheet_state", "hidden"),
            lambda workbook: setattr(workbook, "security", WorkbookProtection(lockStructure=True)),
        )
        for index, change in enumerate(changes):
            with self.subTest(index=index):
                workbook = build_workbook()
                change(workbook)
                self.assert_issue(self.parse(workbook), code="INVALID_WORKBOOK")

    def test_unknown_or_missing_required_sheet_is_rejected(self):
        workbook = build_workbook()
        workbook.create_sheet("其他数据")
        self.assert_issue(self.parse(workbook), code="INVALID_WORKBOOK")
        workbook = build_workbook()
        workbook.remove(workbook["应用文"])
        self.assert_issue(self.parse(workbook), code="INVALID_WORKBOOK")

    def test_extension_mismatch_corruption_and_empty_input(self):
        valid = save_workbook(build_workbook())
        for raw, filename in ((valid, "test.xlsm"), (valid, "test.xls"), (b"not a zip", "test.xlsx"),
                              (valid[:100], "test.xlsx"), (b"", "test.xlsx")):
            with self.subTest(filename=filename, size=len(raw)):
                self.assert_issue(parse_workbook(raw, filename), code="INVALID_WORKBOOK")

    def test_upload_and_archive_member_limits_before_openpyxl(self):
        valid = save_workbook(build_workbook())
        oversized = change_package(valid, additions={f"extra/{index}": b"" for index in range(2001)})
        for raw in (b"x" * (MAX_UPLOAD_BYTES + 1), oversized):
            with self.subTest(size=len(raw)), patch("writing.import_files.load_workbook") as loader:
                self.assert_issue(parse_workbook(raw, "test.xlsx"), code="INVALID_WORKBOOK")
                loader.assert_not_called()

    def test_compressed_bomb_is_rejected_before_openpyxl(self):
        raw = change_package(save_workbook(build_workbook()), additions={"large.bin": b"0" * (MAX_EXPANDED_BYTES + 1)})
        self.assertLess(len(raw), MAX_UPLOAD_BYTES)
        with patch("writing.import_files.load_workbook") as loader:
            self.assert_issue(parse_workbook(raw, "file.xlsx"), code="INVALID_WORKBOOK")
            loader.assert_not_called()

    def test_duplicate_zip_paths_traversal_and_bad_crc_are_rejected(self):
        valid = save_workbook(build_workbook())
        for name in ("../escaped.xml", "xl/../escaped.xml", "/escaped.xml", "xl\\escaped.xml"):
            with self.subTest(name=name):
                self.assert_issue(parse_workbook(change_package(valid, additions={name: b"data"}), "test.xlsx"), code="INVALID_WORKBOOK")
        output = BytesIO()
        with ZipFile(output, "w", ZIP_DEFLATED) as archive:
            archive.writestr("[Content_Types].xml", b"first")
            with self.assertWarns(UserWarning):
                archive.writestr("[Content_Types].xml", b"second")
        self.assert_issue(parse_workbook(output.getvalue(), "test.xlsx"), code="INVALID_WORKBOOK")
        # Damage compressed worksheet bytes, preserving a structurally readable ZIP.
        damaged = bytearray(valid)
        with ZipFile(BytesIO(valid)) as archive:
            info = archive.getinfo("xl/worksheets/sheet1.xml")
            offset = info.header_offset + 30 + len(info.filename.encode()) + len(info.extra)
        damaged[offset + 10] ^= 0xFF
        self.assert_issue(parse_workbook(bytes(damaged), "test.xlsx"), code="INVALID_WORKBOOK")

    def test_xml_entities_utf16_entities_macros_media_and_external_links_are_rejected(self):
        valid = save_workbook(build_workbook())
        entities = b'<!DOCTYPE workbook [<!ENTITY x "EXPAND">]><workbook>&x;</workbook>'
        for name, payload in (("xl/media/image1.png", b"png"), ("xl/vbaProject.bin", b"macro"),
                              ("xl/embeddings/oleObject1.bin", b"object"), ("extra.xml", entities),
                              ("extra.xml", entities.decode().encode("utf-16"))):
            with self.subTest(name=name, size=len(payload)), patch("writing.import_files.load_workbook") as loader:
                self.assert_issue(parse_workbook(change_package(valid, additions={name: payload}), "file.xlsx"), code="INVALID_WORKBOOK")
                loader.assert_not_called()
        raw = change_package(valid, {"xl/_rels/workbook.xml.rels": lambda value: value.replace(
            b"</Relationships>", b'<Relationship Id="external" Type="externalLink" Target="https://example.com/data.xlsx" TargetMode="External"/></Relationships>')})
        self.assert_issue(parse_workbook(raw, "file.xlsx"), code="INVALID_WORKBOOK")
        macro = change_package(valid, {"[Content_Types].xml": lambda value: value.replace(
            b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml", b"application/vnd.ms-excel.sheet.macroEnabled.main+xml")})
        self.assert_issue(parse_workbook(macro, "file.xlsx"), code="INVALID_WORKBOOK")

    def test_external_hyperlink_is_rejected(self):
        workbook = build_workbook()
        workbook["应用文"]["F2"].hyperlink = "https://example.com/"
        self.assert_issue(self.parse(workbook), code="INVALID_WORKBOOK")


class WorkbookDownloadTests(SimpleTestCase):
    def test_generated_template_has_empty_data_tables_and_textual_guidance(self):
        raw = make_template(
            categories=[{"id": "daily-writing", "name": "=分类名称", "sort_order": 10}],
            genres=[{"id": "notice", "name": "通知", "sort_order": 20, "is_active": False}],
        )
        workbook = load_workbook(BytesIO(raw), data_only=False)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ["填写说明", "应用文", "分类", "文种"])
        for title, headers in (("应用文", CONTENT_HEADERS), ("分类", CATEGORY_HEADERS), ("文种", GENRE_HEADERS)):
            self.assertEqual(tuple(cell.value for cell in workbook[title][1]), headers)
            self.assertFalse(any(cell.value is not None for row in workbook[title].iter_rows(min_row=2) for cell in row))
        texts = []
        for row in workbook["填写说明"]:
            for cell in row:
                if cell.value is not None:
                    self.assertEqual(cell.data_type, "s")
                    texts.append(cell.value)
        self.assertIn("=分类名称", texts)
        self.assertIn("停用（不可引用）", texts)
        for kind in ("写作模板", "范文示例", "历年真题", "专项练习"):
            self.assertIn(kind, texts)
        result = parse_workbook(raw, "template.xlsx")
        self.assertEqual(result["contents"], [])
        self.assertEqual(result["categories"], [])
        self.assertEqual(result["genres"], [])
        self.assertEqual([issue["code"] for issue in result["issues"]], ["REQUIRED_FIELD"])

    def test_committed_template_is_compatible_and_never_imports_examples(self):
        root = Path(__file__).resolve().parents[3]
        template = root / "outputs" / "01a157bb-432c-7b63-8e01-27b4024955bf" / "应用文批量导入模板-v1.xlsx"
        self.assertTrue(template.exists(), "已交付模板应作为兼容性基线保留。")
        raw = template.read_bytes()
        result = parse_workbook(raw, template.name)
        self.assertEqual([issue["code"] for issue in result["issues"]], ["REQUIRED_FIELD"])
        self.assertEqual(result["contents"], [])
        workbook = load_workbook(BytesIO(raw))
        for index, value in enumerate(content_row(), 1):
            workbook["应用文"].cell(2, index).value = value
        result = parse_workbook(save_workbook(workbook), template.name)
        self.assertEqual(result["issues"], [])
        self.assertEqual(len(result["contents"]), 1)

    def test_error_report_forces_every_cell_to_text_and_preserves_positions(self):
        issues = [{"sheet": "应用文", "row": 7, "column": "正文", "code": "INVALID_VALUE", "message": value}
                  for value in ("=1+1", "+SUM(A1:A2)", "-1+1", "@SUM(A1)", "中文建议\n第二行")]
        workbook = load_workbook(BytesIO(make_error_report(issues)), data_only=False)
        self.addCleanup(workbook.close)
        worksheet = workbook["错误报告"]
        self.assertEqual(worksheet["B2"].value, "7")
        for row, issue in enumerate(issues, 2):
            self.assertEqual(worksheet.cell(row, 5).value, issue["message"])
        for row in worksheet:
            for cell in row:
                self.assertEqual(cell.data_type, "s")

    def test_empty_error_report_has_only_a_header(self):
        workbook = load_workbook(BytesIO(make_error_report([])))
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.active.max_row, 1)
