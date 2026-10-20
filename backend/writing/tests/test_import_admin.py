from datetime import timedelta
from io import BytesIO
from tempfile import TemporaryDirectory
from unittest.mock import patch
from uuid import uuid4

from django.contrib.admin.models import LogEntry
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, RequestFactory, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from config.admin_navigation import navigation
from distribution.models import AgentProfile
from operations.admin_dashboard import dashboard_callback
from writing.models import WritingCategory, WritingContent, WritingRevision


class WritingImportAdminTests(TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory(prefix="writing-admin-test-")
        self.addCleanup(self.directory.cleanup)
        self.settings_override = override_settings(WRITING_IMPORT_TMP_DIR=self.directory.name)
        self.settings_override.enable()
        self.addCleanup(self.settings_override.disable)
        self.user = get_user_model().objects.create_superuser(
            username="writing-import-admin", password="test-password",
        )
        self.client.force_login(self.user)
        self.category = WritingCategory.objects.create(id="daily", name="日常应用文")
        self.import_url = reverse("admin:writing_writingcontent_import")
        self.confirm_url = reverse("admin:writing_writingcontent_import_confirm")
        self.template_url = reverse("admin:writing_writingcontent_import_template")
        self.list_url = reverse("admin:writing_writingcontent_changelist")

    def workbook(self, *, content_id="notice-001", title="通知写作格式", body="通知\n  保留缩进\n<script>alert(1)</script>", category="daily", rows=None):
        book = Workbook()
        sheet = book.active
        sheet.title = "应用文"
        sheet.append(["内容编号", "分类编号", "文种编号", "内容类型", "标题", "正文", "参考答案", "排序"])
        for row in rows or [[content_id, category, "", "写作模板", title, body, "", 0]]:
            sheet.append(row)
        output = BytesIO()
        book.save(output)
        return SimpleUploadedFile("应用文.xlsx", output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def preview(self, **kwargs):
        return self.client.post(self.import_url, {"import_file": self.workbook(**kwargs)})

    def confirm(self, preview, client=None):
        return (client or self.client).post(self.confirm_url, {
            "preview_id": preview["previewId"], "confirmation_token": preview["confirmationToken"],
        })

    def staff(self, name, *permissions):
        user = get_user_model().objects.create_user(username=name, password="test-password", is_staff=True)
        for value in permissions:
            app, codename = value.split(".")
            user.user_permissions.add(Permission.objects.get(content_type__app_label=app, codename=codename))
        return user

    def test_list_sidebar_dashboard_and_default_mode(self):
        response = self.client.get(self.list_url)
        self.assertContains(response, "下载模板")
        self.assertContains(response, "批量导入")
        self.assertContains(response, "应用文导入")
        response = self.client.get(self.import_url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["form"].fields["mode"].initial, "create_only")
        self.assertContains(response, 'enctype="multipart/form-data"')
        self.assertContains(response, "共享首屏免费集合")
        request = RequestFactory().get("/admin/")
        request.user = self.user
        links = dashboard_callback(request, {})["workbench_links"]
        self.assertIn({"title": "导入应用文", "url": self.import_url}, links)

    def test_template_is_a_real_xlsx_with_empty_business_sheets(self):
        response = self.client.get(self.template_url)
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertIn("no-store", response["Cache-Control"])
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(set(workbook.sheetnames), {"应用文", "分类", "文种", "填写说明"})
        self.assertEqual(workbook["应用文"].max_row, 1)
        self.assertIn("daily", str(list(workbook["填写说明"].values)))

    def test_preview_is_read_only_and_escapes_text(self):
        response = self.preview()
        self.assertEqual(response.status_code, 200)
        preview = response.context["preview"]
        self.assertEqual(preview["mode"], "create_only")
        self.assertEqual(preview["counts"]["create"], 1)
        self.assertTrue(preview["canConfirm"])
        self.assertContains(response, "确认导入草稿")
        self.assertContains(response, "&lt;script&gt;alert(1)&lt;/script&gt;")
        self.assertNotContains(response, "<script>alert(1)</script>")
        self.assertEqual(preview["rows"][0]["values"]["body"], "通知\n  保留缩进\n<script>alert(1)</script>")
        self.assertFalse(WritingContent.objects.exists())

    def test_confirm_redirects_to_refreshable_result_and_scoped_list(self):
        other = WritingContent.objects.create(id="outside", category=self.category, kind="template", title="其他内容", body="其他正文")
        response = self.preview()
        preview = response.context["preview"]
        result_response = self.confirm(preview)
        self.assertEqual(result_response.status_code, 302)
        result = self.client.get(result_response.url)
        self.assertContains(result, "草稿导入完成")
        self.assertContains(result, "查看本批草稿")
        self.assertContains(result, reverse("admin:writing_writingcontent_preview", args=("notice-001",)))
        content = WritingContent.objects.get(pk="notice-001")
        self.assertFalse(content.is_published)
        self.assertIsNone(content.published_revision_id)
        batch = self.client.get(result.context["batch_list_url"])
        self.assertEqual(batch.status_code, 200)
        self.assertEqual(list(batch.context["cl"].queryset.values_list("pk", flat=True)), [content.pk])
        self.assertNotIn(other.pk, list(batch.context["cl"].queryset.values_list("pk", flat=True)))
        self.assertContains(batch, "正在查看本批草稿")

    def test_confirm_replay_does_not_rewrite_after_later_edit(self):
        preview = self.preview().context["preview"]
        self.assertEqual(self.confirm(preview).status_code, 302)
        count = LogEntry.objects.count()
        WritingContent.objects.filter(pk="notice-001").update(title="之后的修改")
        self.assertEqual(self.confirm(preview).status_code, 302)
        self.assertEqual(WritingContent.objects.get(pk="notice-001").title, "之后的修改")
        self.assertEqual(LogEntry.objects.count(), count)

    def test_upsert_shows_changes_and_preserves_published_revision(self):
        content = WritingContent.objects.create(id="notice-001", category=self.category, kind="template", title="旧标题", body="旧正文", sort_order=7)
        revision = WritingRevision.objects.create(
            content=content, version=1, category=self.category, kind="template", title=content.title, body=content.body, sort_order=7, created_by=self.user,
        )
        content.is_published = True
        content.published_revision = revision
        content.save()
        response = self.client.post(self.import_url, {"import_file": self.workbook(), "mode": "upsert_draft"})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "旧标题")
        self.assertContains(response, "导入前")
        self.assertContains(response, "导入后")
        self.assertEqual(response.context["preview"]["counts"]["update"], 1)
        self.assertEqual(self.confirm(response.context["preview"]).status_code, 302)
        content.refresh_from_db()
        revision.refresh_from_db()
        self.assertEqual(content.title, "通知写作格式")
        self.assertEqual(content.sort_order, 0)
        self.assertTrue(content.is_published)
        self.assertEqual(content.published_revision_id, revision.pk)
        self.assertEqual(revision.title, "旧标题")

    def test_invalid_rows_disable_confirmation_and_provide_private_report(self):
        response = self.preview(category="missing")
        self.assertEqual(response.status_code, 400)
        self.assertContains(response, "有错误，暂不能确认", status_code=400)
        self.assertNotContains(response, 'name="confirmation_token"', status_code=400)
        preview = response.context["preview"]
        self.assertFalse(preview["canConfirm"])
        self.assertIsNone(preview.get("confirmationToken"))
        report = self.client.get(response.context["report_url"])
        self.assertEqual(report.status_code, 200)
        workbook = load_workbook(BytesIO(report.content))
        self.assertIn("RELATED_NOT_FOUND", str(list(workbook.active.values)))
        other = get_user_model().objects.create_superuser(username="other-importer", password="test-password")
        self.client.force_login(other)
        self.assertIn(self.client.get(response.context["report_url"]).status_code, (403, 404))
        self.assertIn(self.client.get(self.import_url, {"preview_id": preview["previewId"]}).status_code, (403, 404))
        self.assertFalse(WritingContent.objects.exists())

    def test_missing_file_invalid_mode_and_bad_workbook_return_400(self):
        self.assertEqual(self.client.post(self.import_url, {}).status_code, 400)
        self.assertEqual(self.client.post(self.import_url, {"import_file": self.workbook(), "mode": "delete_all"}).status_code, 400)
        response = self.client.post(self.import_url, {"import_file": SimpleUploadedFile("bad.xlsx", b"not an xlsx")})
        self.assertEqual(response.status_code, 400)
        self.assertFalse(WritingContent.objects.exists())

    def test_upload_larger_than_five_mib_is_rejected_before_preview(self):
        response = self.client.post(self.import_url, {
            "import_file": SimpleUploadedFile("too-large.xlsx", b"0" * (5 * 1024 * 1024 + 1)),
        })
        self.assertEqual(response.status_code, 400)
        self.assertContains(response, "文件超过 5 MiB", status_code=400)
        self.assertNotIn("preview", response.context_data)
        self.assertFalse(WritingContent.objects.exists())

    def test_confirm_requires_signed_fields_and_rejects_tampering(self):
        self.assertEqual(self.client.post(self.confirm_url, {}).status_code, 400)
        preview = self.preview().context["preview"]
        response = self.client.post(self.confirm_url, {
            "preview_id": preview["previewId"], "confirmation_token": preview["confirmationToken"] + "tampered",
        })
        self.assertIn(response.status_code, (400, 403))
        self.assertFalse(WritingContent.objects.exists())

    def test_conflict_expired_and_missing_preview_statuses(self):
        preview = self.preview().context["preview"]
        WritingCategory.objects.filter(pk=self.category.pk).update(name="预览后改名")
        response = self.confirm(preview)
        self.assertEqual(response.status_code, 409)
        self.assertFalse(WritingContent.objects.exists())
        future = timezone.now() + timedelta(minutes=31)
        with patch("writing.import_storage.timezone.now", return_value=future):
            self.assertEqual(self.client.get(self.import_url, {"preview_id": preview["previewId"]}).status_code, 410)
        self.assertEqual(self.client.get(self.import_url, {"preview_id": str(uuid4())}).status_code, 404)

    def test_service_unavailable_returns_chinese_503_without_exception_details(self):
        with patch("writing.import_admin.save_preview", side_effect=OSError("/private/server/path")):
            response = self.preview()
        self.assertEqual(response.status_code, 503)
        self.assertContains(response, "导入服务暂时不可用", status_code=503)
        self.assertNotContains(response, "/private/server/path", status_code=503)

    def test_uncertain_confirmation_keeps_a_link_to_original_batch(self):
        preview = self.preview().context["preview"]
        with patch("writing.import_admin.confirm_import", side_effect=OSError("temporary storage failure")):
            response = self.confirm(preview)
        self.assertEqual(response.status_code, 503)
        self.assertContains(response, "核对原批次结果", status_code=503)
        self.assertIn(preview["previewId"], response.context_data["retry_preview_url"])

    def test_read_only_publish_only_and_unprivileged_staff_are_denied_at_every_stage(self):
        for name, permissions in (
            ("read-only", ("writing.view_writingcontent",)),
            ("publisher", ("writing.view_writingcontent", "writing.publish_writingcontent")),
            ("no-writing", ()),
        ):
            with self.subTest(name=name):
                user = self.staff(name, *permissions)
                self.client.force_login(user)
                self.assertEqual(self.client.get(self.import_url).status_code, 403)
                self.assertEqual(self.client.post(self.import_url, {"import_file": self.workbook()}).status_code, 403)
                self.assertEqual(self.client.post(self.confirm_url, {}).status_code, 403)
                self.assertEqual(self.client.get(self.template_url).status_code, 403)
                self.assertEqual(self.client.get(reverse("admin:writing_writingcontent_import_report", args=(uuid4(),))).status_code, 403)
                request = RequestFactory().get("/admin/")
                request.user = user
                self.assertNotIn("应用文导入", [item["title"] for group in navigation(request) for item in group["items"]])
                self.assertNotIn("导入应用文", [item["title"] for item in dashboard_callback(request, {})["workbench_links"]])

    def test_editor_without_publish_permission_can_import_but_needs_change_for_upsert(self):
        editor = self.staff("creator", "writing.view_writingcontent", "writing.add_writingcontent")
        self.client.force_login(editor)
        self.assertEqual(self.client.get(self.import_url).status_code, 200)
        preview = self.preview().context["preview"]
        self.assertEqual(self.confirm(preview).status_code, 302)
        response = self.client.post(self.import_url, {"import_file": self.workbook(title="需修改"), "mode": "upsert_draft"})
        self.assertIn(response.status_code, (400, 403))
        if response.context_data.get("preview"):
            self.assertFalse(response.context_data["preview"]["canConfirm"])
        self.assertEqual(WritingContent.objects.get(pk="notice-001").title, "通知写作格式")

    def test_permission_revocation_blocks_confirm_report_and_batch_list(self):
        editor = self.staff("revoked-editor", "writing.view_writingcontent", "writing.add_writingcontent")
        self.client.force_login(editor)
        response = self.preview()
        preview = response.context["preview"]
        editor.user_permissions.clear()
        self.assertEqual(self.confirm(preview).status_code, 403)
        self.assertEqual(self.client.get(response.context["report_url"]).status_code, 403)
        self.assertEqual(self.client.get(self.list_url, {"import_batch": preview["previewId"]}).status_code, 403)
        self.assertFalse(WritingContent.objects.exists())

    def test_agent_self_account_is_denied_even_with_content_permissions(self):
        user = self.staff(
            "agent-importer", "writing.view_writingcontent", "writing.add_writingcontent",
            "distribution.view_agentprofile", "distribution.view_agentorderattribution",
            "distribution.view_agentvoucher", "distribution.view_agentperformance",
            "distribution.export_agentperformance", "activation.view_activationredemption",
        )
        AgentProfile(user=user, agent_code="WRITING-AGENT", display_name="代理本人", created_by=self.user).save(_service=True)
        self.client.force_login(user)
        self.assertEqual(self.client.get(self.import_url).status_code, 403)
        self.assertEqual(self.client.post(self.import_url, {"import_file": self.workbook()}).status_code, 403)
        self.assertEqual(self.client.post(self.confirm_url).status_code, 403)

    def test_anonymous_redirects_to_login_and_unsafe_methods_are_rejected(self):
        self.client.logout()
        for url in (self.import_url, self.template_url, reverse("admin:writing_writingcontent_import_report", args=(uuid4(),))):
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertEqual(response.status_code, 302)
                self.assertTrue(response.url.startswith(reverse("admin:login")))
        self.assertEqual(self.client.post(self.confirm_url).status_code, 302)
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(self.confirm_url).status_code, 405)
        self.assertEqual(self.client.post(self.template_url).status_code, 405)
        self.assertEqual(self.client.put(self.import_url).status_code, 405)

    def test_upload_and_confirmation_require_csrf(self):
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.user)
        self.assertEqual(client.post(self.import_url, {"import_file": self.workbook()}).status_code, 403)
        preview = self.preview().context["preview"]
        self.assertEqual(self.confirm(preview, client=client).status_code, 403)
        client.get(self.import_url)
        token = client.cookies["csrftoken"].value
        response = client.post(self.confirm_url, {
            "preview_id": preview["previewId"], "confirmation_token": preview["confirmationToken"], "csrfmiddlewaretoken": token,
        })
        self.assertEqual(response.status_code, 302)
