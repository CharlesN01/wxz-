"""Bounded XLSX parsing and downloads for applied-writing imports.

This module deliberately has no database or Django dependencies. Uploaded cells
are untrusted input; generated text cells are never interpreted as formulas.
"""

import posixpath
import re
from io import BytesIO
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile
from zlib import error as ZlibError

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter


MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_ZIP_MEMBERS = 2000
MAX_CONTENT_ROWS = 1000
MAX_DICTIONARY_ROWS = 200
MAX_SORT_ORDER = 2147483647
MAX_ISSUES = 2000

CONTENT_HEADERS = (
    "内容编号", "分类编号", "文种编号", "内容类型", "标题", "正文", "参考答案", "排序",
)
CATEGORY_HEADERS = ("分类编号", "分类名称", "排序")
GENRE_HEADERS = ("文种编号", "文种名称", "排序")
SHEETS = {
    "应用文": ("contents", CONTENT_HEADERS, MAX_CONTENT_ROWS),
    "分类": ("categories", CATEGORY_HEADERS, MAX_DICTIONARY_ROWS),
    "文种": ("genres", GENRE_HEADERS, MAX_DICTIONARY_ROWS),
}
KIND_LABELS = {
    "写作模板": "template", "范文示例": "example", "历年真题": "past_exam", "专项练习": "exercise",
}
KINDS = {**KIND_LABELS, **{value: value for value in KIND_LABELS.values()}}
ID_PATTERN = re.compile(r"[A-Za-z0-9_-]{1,64}\Z", re.ASCII)
CELL_PATTERN = re.compile(r"([A-Z]{1,3})([1-9][0-9]{0,6})\Z", re.ASCII)
MAIN_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"


class _RejectedWorkbook(ValueError):
    pass


def _issue(issues, sheet, row, column, code, message):
    if len(issues) >= MAX_ISSUES:
        raise _RejectedWorkbook("错误数量过多，请先修正已列出的错误后重新上传。")
    issues.append({"sheet": sheet, "row": row, "column": column, "code": code, "message": message})


def _local(tag):
    return tag.rsplit("}", 1)[-1]


def _truthy(value):
    return value in ("1", "true", "True")


def _xml_events(archive, name):
    with archive.open(name) as stream:
        yield from ElementTree.iterparse(stream, events=("start", "end"))


def _inspect_package(archive):
    """Check the entire ZIP before openpyxl reads styles or shared strings.

    Read every member to validate CRCs and actual expanded size. Checking XML
    declarations before parsing also rejects entities in UTF-16/32 documents.
    """
    members = archive.infolist()
    if len(members) > MAX_ZIP_MEMBERS:
        raise _RejectedWorkbook("工作簿 ZIP 成员超过 2,000 个，请使用标准模板重新保存。")
    names = set()
    expanded = 0
    for member in members:
        name = member.filename
        if (
            name in names or "\\" in name or name.startswith("/")
            or any(part in ("..", ".") for part in name.split("/"))
            or "\x00" in name or ":" in name or member.flag_bits & 1
        ):
            raise _RejectedWorkbook("工作簿 ZIP 路径重复、无效或已加密，请重新保存为未加密的 XLSX。")
        names.add(name)
        expanded += member.file_size
        if expanded > MAX_EXPANDED_BYTES:
            raise _RejectedWorkbook("工作簿解压总大小超过 50 MiB，请减少内容后重试。")
    required = {"[Content_Types].xml", "_rels/.rels", "xl/workbook.xml", "xl/_rels/workbook.xml.rels"}
    if not required.issubset(names):
        raise _RejectedWorkbook("文件不是完整的 XLSX 工作簿，请使用标准模板重新保存。")
    forbidden = (
        "xl/media/", "xl/drawings/", "xl/charts/", "xl/embeddings/", "xl/activex/",
        "xl/externallinks/", "xl/connections", "xl/querytables/", "customui/",
    )
    actual_expanded = 0
    for member in members:
        name = member.filename
        lowered = name.lower()
        if lowered.startswith(forbidden) or "vbaproject" in lowered or "vbasignature" in lowered:
            raise _RejectedWorkbook("工作簿含图片、图表、外部数据、嵌入对象或宏，请删除后重新上传。")
        tail = b""
        is_xml = lowered.endswith((".xml", ".rels"))
        with archive.open(member) as stream:
            while chunk := stream.read(64 * 1024):
                actual_expanded += len(chunk)
                if actual_expanded > MAX_EXPANDED_BYTES:
                    raise _RejectedWorkbook("工作簿实际解压总大小超过 50 MiB，请减少内容后重试。")
                if is_xml:
                    text = tail + chunk.replace(b"\x00", b"").lower()
                    if b"<!doctype" in text or b"<!entity" in text:
                        raise _RejectedWorkbook("工作簿含不支持的 XML 实体声明，请使用标准模板重新保存。")
                    tail = text[-32:]

    sheet_ids = {}
    relations = {}
    workbook_content_type = None
    default_xml_content_type = None
    for name in names:
        if not name.lower().endswith((".xml", ".rels")):
            continue
        for event, element in _xml_events(archive, name):
            tag = _local(element.tag)
            if event == "start":
                if tag == "Relationship":
                    relation_type = element.get("Type", "").lower()
                    target = element.get("Target", "")
                    if element.get("TargetMode", "").lower() == "external" or ":" in target or target.startswith("//"):
                        raise _RejectedWorkbook("工作簿含外部链接，请删除外部链接后重新上传。")
                    if any(part in relation_type for part in ("vbaproject", "oleobject", "activex", "externallink")):
                        raise _RejectedWorkbook("工作簿含宏、外链或嵌入对象，请删除后重新上传。")
                    if name == "xl/_rels/workbook.xml.rels":
                        relations[element.get("Id")] = posixpath.normpath(
                            target.lstrip("/") if target.startswith("/") else posixpath.join("xl", target)
                        )
                elif name == "[Content_Types].xml" and tag in ("Override", "Default"):
                    content_type = element.get("ContentType", "").lower()
                    if any(part in content_type for part in ("macroenabled", "vbaproject", "macrosheet", "activex", "oleobject")):
                        raise _RejectedWorkbook("不支持含宏的工作簿，请保存为不含宏的 XLSX。")
                    if element.get("PartName") == "/xl/workbook.xml":
                        workbook_content_type = content_type
                    elif tag == "Default" and element.get("Extension", "").lower() == "xml":
                        default_xml_content_type = content_type
                elif name == "xl/workbook.xml":
                    if tag == "workbookProtection" and any(
                        _truthy(value) if key in ("lockStructure", "lockWindows", "lockRevision") else bool(value)
                        for key, value in element.attrib.items()
                    ):
                        raise _RejectedWorkbook("工作簿已设置保护，请取消保护后重新上传。")
                    if tag == "externalReferences":
                        raise _RejectedWorkbook("工作簿含外部链接，请删除外部链接后重新上传。")
                    if tag == "sheet":
                        title = element.get("name", "")
                        if title not in (*SHEETS, "填写说明"):
                            raise _RejectedWorkbook("存在未知工作表，请只保留应用文、分类、文种和填写说明。")
                        if title in sheet_ids:
                            raise _RejectedWorkbook("工作表名称重复，请使用标准模板重新保存。")
                        if element.get("state", "visible") != "visible":
                            raise _RejectedWorkbook("不允许隐藏工作表，请取消隐藏后重新上传。")
                        sheet_ids[title] = next((v for k, v in element.attrib.items() if _local(k) == "id"), None)
            else:
                element.clear()
    if (workbook_content_type or default_xml_content_type) != MAIN_CONTENT_TYPE:
        raise _RejectedWorkbook("文件内容不是标准 XLSX，请勿仅修改文件后缀。")
    if "应用文" not in sheet_ids:
        raise _RejectedWorkbook("缺少“应用文”工作表，请使用标准模板填写。")
    sheet_parts = {}
    for title, relation_id in sheet_ids.items():
        part = relations.get(relation_id)
        if not part or part not in names or part in sheet_parts:
            raise _RejectedWorkbook("工作表关联损坏，请使用标准模板重新保存。")
        sheet_parts[part] = title
    return sheet_parts


def _inspect_sheet(archive, part, title, issues):
    """Stream physical rows rather than trusting the worksheet dimension."""
    max_row = max_column = 0
    previous_row = 0
    current_row = None
    current_column = None
    previous_column = 0
    has_value = False
    report_sheet = title if title in SHEETS else "工作簿"
    for event, element in _xml_events(archive, part):
        tag = _local(element.tag)
        if event == "start":
            if tag == "row":
                current_row = int(element.get("r", previous_row + 1))
                if not previous_row < current_row <= 1048576:
                    raise _RejectedWorkbook("工作簿行号无效或重复，请使用标准模板重新保存。")
                previous_row = current_row
                previous_column = 0
                if _truthy(element.get("hidden")):
                    _issue(issues, report_sheet, current_row, None, "INVALID_WORKBOOK", f"“{title}”存在隐藏行，请取消隐藏后重试。")
            elif tag == "col" and _truthy(element.get("hidden")):
                _issue(issues, report_sheet, None, None, "INVALID_WORKBOOK", f"“{title}”存在隐藏列，请取消隐藏后重试。")
            elif tag in ("mergeCells", "sheetProtection", "protectedRanges", "drawing", "legacyDrawing", "picture", "oleObjects", "controls"):
                _issue(issues, report_sheet, None, None, "INVALID_WORKBOOK", f"“{title}”含合并单元格、保护、图片或嵌入对象，请删除后重试。")
            elif tag == "c":
                coordinate = element.get("r")
                if coordinate:
                    match = CELL_PATTERN.fullmatch(coordinate)
                    if not match or int(match[2]) != current_row:
                        raise _RejectedWorkbook("工作簿单元格位置无效，请使用标准模板重新保存。")
                    current_column = column_index_from_string(match[1])
                else:
                    current_column = previous_column + 1
                if current_row is None or not previous_column < current_column <= 16384:
                    raise _RejectedWorkbook("工作簿单元格列号无效或重复，请重新保存。")
                previous_column = current_column
                has_value = False
                if element.get("t") == "e":
                    _issue(issues, report_sheet, current_row, get_column_letter(current_column), "INVALID_VALUE", f"“{title}”存在 Excel 错误值，请改为正确的纯文本或整数。")
            elif tag == "f":
                _issue(issues, report_sheet, current_row, get_column_letter(current_column) if current_column else None, "INVALID_VALUE", f"“{title}”不允许公式，请粘贴为纯文本或整数值。")
        else:
            if tag in ("v", "t") and element.text is not None:
                has_value = True
            elif tag == "c":
                if has_value:
                    max_row = max(max_row, current_row)
                    max_column = max(max_column, current_column)
                current_column = None
            elif tag == "row":
                current_row = None
            element.clear()
    return max_row, max_column


def _empty(value):
    return value is None or isinstance(value, str) and not value.strip()


def _read_text(value, *, issues, sheet, row, column, required=False, maximum=None, strip=False):
    if _empty(value):
        if required:
            _issue(issues, sheet, row, column, "REQUIRED_FIELD", f"{column}不能为空，请填写有效文本。")
        return ""
    if not isinstance(value, str):
        _issue(issues, sheet, row, column, "INVALID_VALUE", f"{column}必须使用文本单元格，请设置为文本后重新填写。")
        return ""
    value = value.strip() if strip else value.replace("\r\n", "\n").replace("\r", "\n")
    if maximum is not None and len(value) > maximum:
        _issue(issues, sheet, row, column, "INVALID_VALUE", f"{column}不能超过 {maximum:,} 个字符，请缩短后重试。")
    return value


def _read_id(value, issues, sheet, row, column, *, required=True):
    value = _read_text(value, issues=issues, sheet=sheet, row=row, column=column, required=required, strip=True)
    if value and not ID_PATTERN.fullmatch(value):
        _issue(issues, sheet, row, column, "INVALID_VALUE", f"{column}须为 1～64 位英文字母、数字、下划线或连字符，请检查编号。")
    return value


def _read_sort(value, issues, sheet, row):
    if value is None or value == "":
        return 0
    if type(value) is not int or not 0 <= value <= MAX_SORT_ORDER:
        _issue(issues, sheet, row, "排序", "INVALID_VALUE", "排序须为 0～2,147,483,647 的真正整数；请勿填写文本、布尔值、小数或日期。")
        return 0
    return value


def _parse_row(data, sheet, row, issues):
    if sheet != "应用文":
        return {
            "id": _read_id(data[f"{sheet}编号"], issues, sheet, row, f"{sheet}编号"),
            "name": _read_text(data[f"{sheet}名称"], issues=issues, sheet=sheet, row=row, column=f"{sheet}名称", required=True, maximum=120, strip=True),
            "sort_order": _read_sort(data["排序"], issues, sheet, row),
        }
    values = {
        "id": _read_id(data["内容编号"], issues, sheet, row, "内容编号"),
        "category_id": _read_id(data["分类编号"], issues, sheet, row, "分类编号"),
        "genre_id": _read_id(data["文种编号"], issues, sheet, row, "文种编号", required=False) or None,
        "title": _read_text(data["标题"], issues=issues, sheet=sheet, row=row, column="标题", required=True, maximum=200, strip=True),
        "body": _read_text(data["正文"], issues=issues, sheet=sheet, row=row, column="正文", required=True, maximum=20000),
        "reference_answer": _read_text(data["参考答案"], issues=issues, sheet=sheet, row=row, column="参考答案", maximum=20000),
        "sort_order": _read_sort(data["排序"], issues, sheet, row),
    }
    kind = _read_text(data["内容类型"], issues=issues, sheet=sheet, row=row, column="内容类型", required=True, strip=True)
    values["kind"] = KINDS.get(kind, "")
    if kind and not values["kind"]:
        _issue(issues, sheet, row, "内容类型", "INVALID_VALUE", "内容类型只能填写写作模板、范文示例、历年真题、专项练习或其英文代码。")
    if values["kind"] == "exercise" and not values["reference_answer"]:
        _issue(issues, sheet, row, "参考答案", "REQUIRED_FIELD", "专项练习必须填写参考答案，请补全后重试。")
    elif values["kind"] and values["kind"] != "exercise" and values["reference_answer"]:
        _issue(issues, sheet, row, "参考答案", "INVALID_VALUE", "只有专项练习可以填写参考答案，请将本行参考答案清空。")
    return values


def _parse_sheet(worksheet, sheet, bounds, result):
    key, headers, limit = SHEETS[sheet]
    issues = result["issues"]
    max_row, max_column = bounds
    worksheet.reset_dimensions()
    columns = max(len(headers), max_column)
    cells = next(worksheet.iter_rows(min_row=1, max_row=1, max_col=columns))
    actual_headers = [cell.value for cell in cells]
    while actual_headers and actual_headers[-1] is None:
        actual_headers.pop()
    if max_column > len(headers) or len(actual_headers) != len(headers) or any(actual_headers.count(header) != 1 for header in headers):
        _issue(issues, sheet, 1, None, "INVALID_HEADER", f"表头必须且只能包含：{'、'.join(headers)}；请补齐缺失列并删除重复、未知或空白列名。")
        return
    seen_ids = {}
    count = 0
    for cells in worksheet.iter_rows(min_row=2, max_row=max(max_row, 1), max_col=len(headers)):
        if all(_empty(cell.value) for cell in cells):
            continue
        row = next(cell.row for cell in cells if cell.value is not None)
        count += 1
        if count > limit:
            _issue(issues, sheet, row, None, "INVALID_VALUE", f"“{sheet}”最多允许 {limit:,} 条非空数据行，请拆分文件。")
            break
        values = _parse_row(dict(zip(actual_headers, (cell.value for cell in cells))), sheet, row, issues)
        identifier = values["id"]
        if identifier:
            if identifier in seen_ids:
                _issue(issues, sheet, row, headers[0], "DUPLICATE_ID", f"编号与第 {seen_ids[identifier]} 行重复，请为每条数据使用唯一编号。")
            else:
                seen_ids[identifier] = row
        result[key].append({"row": row, "values": values})


def parse_workbook(raw: bytes, filename: str) -> dict:
    """Return normalized rows and actionable issues without writing anything."""
    result = {"contents": [], "categories": [], "genres": [], "issues": []}
    issues = result["issues"]
    workbook = None
    try:
        if not isinstance(filename, str) or not filename.lower().endswith(".xlsx"):
            raise _RejectedWorkbook("仅支持 .xlsx 文件，请下载标准模板后重新填写。")
        if not isinstance(raw, bytes) or not raw:
            raise _RejectedWorkbook("上传文件为空或无效，请重新选择 XLSX 文件。")
        if len(raw) > MAX_UPLOAD_BYTES:
            raise _RejectedWorkbook("上传文件超过 5 MiB，请减少内容后重新上传。")
        with ZipFile(BytesIO(raw)) as archive:
            sheet_parts = _inspect_package(archive)
            bounds = {title: _inspect_sheet(archive, part, title, issues) for part, title in sheet_parts.items()}
        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        # Translate physical column letters into the uploaded (possibly reordered)
        # Chinese headers, without calculating formulas or trusting dimensions.
        for issue in issues:
            if issue["column"] and issue["sheet"] in SHEETS:
                worksheet = workbook[issue["sheet"]]
                index = column_index_from_string(issue["column"])
                header = next(worksheet.iter_rows(min_row=1, max_row=1, min_col=index, max_col=index))[0]
                if isinstance(header.value, str) and header.value in SHEETS[issue["sheet"]][1]:
                    issue["column"] = header.value
        if issues:
            return result
        for sheet in SHEETS:
            if sheet in workbook.sheetnames:
                _parse_sheet(workbook[sheet], sheet, bounds[sheet], result)
        if not result["contents"]:
            _issue(issues, "应用文", None, None, "REQUIRED_FIELD", "应用文工作表没有可导入内容，请从第 2 行开始填写至少一篇应用文。")
    except _RejectedWorkbook as error:
        if len(issues) < MAX_ISSUES:
            _issue(issues, "工作簿", None, None, "INVALID_WORKBOOK", str(error))
    except (BadZipFile, ZlibError, ElementTree.ParseError, OSError, RuntimeError, ValueError, KeyError, IndexError, TypeError, AttributeError, OverflowError, NotImplementedError):
        if len(issues) < MAX_ISSUES:
            _issue(issues, "工作簿", None, None, "INVALID_WORKBOOK", "文件损坏、加密或内部格式无效，请用 Excel 重新保存为标准 XLSX 后上传。")
    finally:
        if workbook is not None:
            workbook.close()
    return result


def _text_cell(worksheet, row, column, value):
    cell = worksheet.cell(row=row, column=column)
    cell.value = "" if value is None else str(value)
    cell.data_type = "s"
    cell.number_format = "@"
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    return cell


def _format_table(worksheet, headers):
    for index, header in enumerate(headers, 1):
        cell = _text_cell(worksheet, 1, index, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305A78")
        worksheet.column_dimensions[get_column_letter(index)].width = 22
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def _workbook_bytes(workbook):
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def make_template(categories=(), genres=()) -> bytes:
    """Generate empty import tables; examples and current IDs live in guidance."""
    workbook = Workbook()
    guidance = workbook.active
    guidance.title = "填写说明"
    _format_table(guidance, ("项目", "填写说明", "补充", "排序"))
    instructions = (
        ("应用文批量导入模板", "版本 v1；三个数据表均留空，样例不属于正式内容。"),
        ("填写顺序", "先查阅下方已有分类、文种编号，再在“应用文”工作表从第 2 行填写。"),
        ("文件限制", "仅支持 XLSX；文件不超过 5 MiB，解压总大小不超过 50 MiB，ZIP 成员不超过 2,000 个。"),
        ("数量限制", "应用文最多 1,000 条，分类、文种各最多 200 条非空数据行。"),
        ("表头", "列名须与模板精确一致；可以换列顺序，不得增删规定列或增加其他工作表。"),
        ("编号", "编号列请先设为文本；1～64 位，仅英文字母、数字、下划线、连字符，保留大小写和前导零。"),
        ("标题及名称", "标题最多 200 字符；分类和文种名称最多 120 字符；均去除首尾空白。"),
        ("正文及答案", "最多各 20,000 字符；Alt+Enter 输入真实换行，保留缩进与段落；仅按纯文本显示。"),
        ("排序", "填写真正的 0～2,147,483,647 整数，留空为 0；越小越靠前，同值按内容编号排序。"),
        ("写作模板 / template", "参考答案留空；模板全文写入正文。"),
        ("范文示例 / example", "参考答案留空；范文全文写入正文。"),
        ("历年真题 / past_exam", "参考答案留空；材料、题干、解析和答案按段落统一写入正文。"),
        ("专项练习 / exercise", "题目写入正文，参考答案必须填写。"),
        ("分类、文种与类型", "分类是内容归属；文种是通知、请假条等体裁；内容类型决定进入四个模块中的哪一个。"),
        ("字典表", "仅填写本次需要新增且被内容引用的字典；已有编号直接引用，不能通过导入修改或启用字典。"),
        ("新增及更新草稿", "完整替换草稿字段；文种留空清除原文种，排序留空重置为 0；文件外记录不变。"),
        ("发布与免费规则", "导入仅保存草稿；发布沿用独立流程。免费集合由现有全类型共享排序及权益规则决定，没有逐篇免费或价格字段。"),
        ("禁止结构", "请删除隐藏行列、隐藏表、合并、公式、Excel 错误值、外链、图片、图表、宏和密码保护。"),
    )
    for row, values in enumerate(instructions, 2):
        for column, value in enumerate(values, 1):
            _text_cell(guidance, row, column, value)
    row = len(instructions) + 3
    examples = (
        ("样例分类", "daily-writing", "日常应用文", "10"),
        ("样例文种", "notice", "通知", "10"),
        ("样例文种", "leave-note", "请假条", "20"),
        ("四类型样例（仅格式示例）", "以下内容不是正式历年试题，不会自动导入。"),
        ("内容编号", "分类编号", "文种编号", "内容类型", "标题", "正文", "参考答案", "排序"),
        ("notice-template-001", "daily-writing", "notice", "写作模板", "通知写作格式", "标题：通知\n称谓：通知对象\n正文：事由、时间、地点、要求\n落款：单位及日期", "", "10"),
        ("leave-example-001", "daily-writing", "leave-note", "范文示例", "请假条范文", "请假条\n尊敬的老师：\n  因身体不适，申请请假一天。\n申请人：张同学", "", "20"),
        ("notice-past-001", "daily-writing", "notice", "历年真题", "通知写作格式示例", "请根据题目要求写一则通知。\n参考要点：对象明确，时间、地点齐全。", "", "30"),
        ("notice-exercise-001", "daily-writing", "notice", "专项练习", "写一则班会通知", "请以班委会名义拟写一则班会通知。", "通知\n全体同学：\n  请于指定时间到教室参加班会。\n班委会", "40"),
    )
    for values in examples:
        for column, value in enumerate(values, 1):
            _text_cell(guidance, row, column, value)
        row += 1
    for title, entries in (("已有分类清单", categories), ("已有文种清单", genres)):
        row += 1
        _text_cell(guidance, row, 1, title)
        row += 1
        for column, value in enumerate(("编号", "名称", "排序", "状态"), 1):
            _text_cell(guidance, row, column, value)
        row += 1
        for entry in entries:
            for column, value in enumerate((entry["id"], entry["name"], entry.get("sort_order", 0), "启用" if entry.get("is_active", True) else "停用（不可引用）"), 1):
                _text_cell(guidance, row, column, value)
            row += 1
    for sheet, (_, headers, limit) in SHEETS.items():
        worksheet = workbook.create_sheet(sheet)
        _format_table(worksheet, headers)
        for index, header in enumerate(headers, 1):
            if "编号" in header:
                # Preformat blank cells so entering 001 does not lose its zeros.
                for data_row in range(2, limit + 2):
                    worksheet.cell(data_row, index).number_format = "@"
            if header in ("正文", "参考答案"):
                worksheet.column_dimensions[get_column_letter(index)].width = 65
    guidance.column_dimensions["A"].width = 27
    guidance.column_dimensions["B"].width = 75
    return _workbook_bytes(workbook)


def make_error_report(issues) -> bytes:
    """Use XLSX string cells even when a message begins with =, +, - or @."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "错误报告"
    _format_table(worksheet, ("工作表", "Excel 原始行号", "列名", "错误码", "中文修改建议"))
    for row, issue in enumerate(issues, 2):
        for column, key in enumerate(("sheet", "row", "column", "code", "message"), 1):
            _text_cell(worksheet, row, column, issue.get(key))
    worksheet.column_dimensions["E"].width = 85
    return _workbook_bytes(workbook)
